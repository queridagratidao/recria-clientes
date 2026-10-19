import html, json
from docx import Document
from docx.shared import Pt, RGBColor
BASE="https://www.agenciarecria.com.br/recria-ads/?utm_source=meta&utm_medium=paid&utm_campaign=recria-ads&utm_content="
C=[
("Criativo 1 · Imagem única com preço","ads-1-dor.png (feed 1x1) e ads-1-dor.png (stories 9:16)","c1-dor",
"Você cria conteúdo e anúncio, mas quase ninguém compra? 🤔\nIsso não é falta de esforço: é que o jeito de vender mudou.\nNo Recria Ads, você aprende a criar conteúdo e anúncio sem cara de anúncio e aumenta as vendas do seu negócio, com aulas de 10 minutos por dia.\n💳 Apenas 12x de R$ 10,03.\n👉 Clique em \"Quero saber mais\"."),
("Criativo 2 · Carrossel de 5 cards","ads-2a-historia, ads-2b-revelacao, ads-2c-porque, ads-2d-gancho, ads-2e-oferta (nesta ordem)","c2-virada",
"Se você tem um negócio, eu preciso te contar uma coisa que pode estar te fazendo perder dinheiro nesse exato momento. Lá atrás, isso que eu vou te contar me fez quebrar dois negócios. Depois que eu descobri e consegui ajustar, tudo mudou… 👀\nArraste para o lado. 👉\nNo Recria Ads, você aprende a criar conteúdo e anúncio sem cara de anúncio e vende com leveza e naturalidade, sem parecer aquele vendedor chato.\n👉 Clique em \"Quero saber mais\"."),
("Criativo 3 · Imagem única com o post \"7 Verdades Duras\"","ads-3-prova.png (feed 1x1 e stories 9:16)","c3-prova",
"Esse carrossel é do meu perfil de yoga. Atraí as pessoas, elas se identificaram e comentaram \"eu quero\". No direct, continuei a conversa com a oferta do meu curso e vendi. Tudo configurado de forma automatizada.\nQuer aprender a criar conteúdos que atraem, engajam e vendem, com leveza? É isso que eu ensino no Recria Ads.\n👉 Clique em \"Quero saber mais\".\nCaso real do meu perfil de yoga. Os resultados variam de pessoa para pessoa."),
("Criativo 4 · Carrossel \"passou na faculdade, ganhou o carro\"","Carrossel da Amanda (já pronto)","c4-carrossel",
"[LEGENDA PENDENTE: falta o texto do carrossel para eu escrever no mesmo estilo]"),
]
CFG=["Objetivo: Vendas, otimizando para Compra (Purchase).","Um conjunto de anúncios só, com os 4 criativos. Brasil, público amplo, sem segmentação de interesse.","Verba de teste: R$ 50 por dia, por 7 dias (R$ 350 no total). Não editar nos primeiros 3 dias.","Depois de 3 dias: desligar o criativo que gastou cerca de R$ 100 sem nenhuma venda e pôr mais verba no que vendeu (regra prática, não é número oficial da Meta).","Antes de ligar: clicar em todos os botões da página de vendas, fazer uma compra de teste até a página de obrigado e conferir no Gerenciador de Eventos: ViewContent, InitiateCheckout e Purchase (sem contar a compra em dobro)."]
T="Recria Ads · Legendas dos criativos"
# DOCX
d=Document()
st=d.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(11)
d.add_heading(T,0)
d.add_paragraph("Para consultar na hora de subir a campanha. Copie a legenda e o link de cada criativo.")
d.add_heading("Configuração sugerida",1)
for c in CFG: d.add_paragraph(c,style='List Bullet')
for t,a,u,l in C:
    d.add_heading(t,1)
    p=d.add_paragraph();r=p.add_run("Arquivos: ");r.bold=True;p.add_run(a)
    p=d.add_paragraph();r=p.add_run("Link de destino: ");r.bold=True;p.add_run(BASE+u)
    p=d.add_paragraph();r=p.add_run("Legenda:");r.bold=True
    for line in l.split("\n"): d.add_paragraph(line)
d.save("Legendas-Recria-Ads.docx")
# HTML -> PDF
h=f"""<!doctype html><meta charset=utf-8><style>
body{{font-family:'Helvetica Neue',Arial,sans-serif;color:#1c1c1c;margin:0;padding:0;font-size:12.5px;line-height:1.5}}
h1{{font-size:24px;margin:0 0 4px}}h2{{font-size:16px;margin:26px 0 6px;border-bottom:2px solid #C9A24E;padding-bottom:4px}}
.s{{color:#666;margin-bottom:14px}}.b{{font-weight:700}}
.leg{{background:#faf6ec;border-left:4px solid #C9A24E;padding:10px 14px;margin:6px 0 10px}}
.leg p{{margin:0 0 6px}}.u{{word-break:break-all;font-size:10.5px;color:#444}}
li{{margin-bottom:3px}}.box{{page-break-inside:avoid}}
</style><h1>{T}</h1><div class=s>Para consultar na hora de subir a campanha. Copie a legenda e o link de cada criativo.</div>
<h2>Configuração sugerida</h2><ul>{''.join('<li>'+html.escape(c)+'</li>' for c in CFG)}</ul>"""
for t,a,u,l in C:
    h+=f"<div class=box><h2>{html.escape(t)}</h2><div><span class=b>Arquivos:</span> {html.escape(a)}</div><div class=u><span class=b>Link de destino:</span> {html.escape(BASE+u)}</div><div class=b style='margin-top:8px'>Legenda:</div><div class=leg>{''.join('<p>'+html.escape(x)+'</p>' for x in l.split(chr(10)))}</div></div>"
open("_t.html","w").write(h)
