#!/usr/bin/env python3
"""Gera a planilha de bônus do e-book Calendário manual. Requer: pip install openpyxl"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

PRETO, OURO, CREME = "0B0B0B", "C9A24E", "F6F0E2"
head_font = Font(bold=True, color="E4C988", name="Calibri", size=11)
head_fill = PatternFill("solid", fgColor=PRETO)
thin = Side(style="thin", color="D9C394")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

def cabecalho(ws, row, cols):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.border = head_font, head_fill, border
        c.alignment = Alignment(wrap_text=True, vertical="center")

def larguras(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w

wb = Workbook()

# --- Leia primeiro
ws = wb.active; ws.title = "Leia primeiro"
linhas = [
 ("PLANILHA DE CALENDÁRIO DE CONTEÚDO · RECRIA", True),
 ("Bônus do e-book Como Criar o Seu Calendário de Conteúdo", False),
 ("", False),
 ("Como usar", True),
 ("1. Aba Cliente: escreva as dores, os desejos, as objeções e as perguntas do seu cliente.", False),
 ("2. Aba Calendário: preencha 1 linha por post. Use as listas de seleção (clique na célula e na setinha).", False),
 ("3. Aba Resumo: veja se as proporções dos pilares estão perto de 40% Educar, 25% Conectar, 20% Provar e 15% Vender.", False),
 ("4. Aba Datas: confira as datas do ano e acrescente as do seu nicho.", False),
 ("5. Aba Banco de ganchos: copie os modelos e troque o que está entre colchetes.", False),
 ("6. Aba Revisão do mês: no fim do mês, anote os resultados por formato para repetir o que funcionou.", False),
 ("", False),
 ("Dicas", True),
 ("- Termine o calendário do mês 1 a 2 semanas antes de o mês começar.", False),
 ("- Nunca coloque dois posts de venda seguidos.", False),
 ("- Escreva o gancho completo de cada post: quem escreve o gancho, publica.", False),
 ("- Confirme as datas comemorativas todo ano: várias são móveis.", False),
 ("", False),
 ("Recria · Recriando marketing, simplificando vendas · agenciarecria.com.br", False),
]
for i, (t, b) in enumerate(linhas, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(bold=b, size=14 if i == 1 else 11, color=PRETO if b else "2B2415")
ws.column_dimensions["A"].width = 120

# --- Cliente
ws = wb.create_sheet("Cliente")
cabecalho(ws, 1, ["Dores", "Desejos", "Objeções", "Perguntas frequentes"])
for r in range(2, 12):
    for c in range(1, 5):
        ws.cell(row=r, column=c).border = border; ws.cell(row=r, column=c).alignment = wrap
larguras(ws, {"A": 40, "B": 40, "C": 40, "D": 40})
ws.freeze_panes = "A2"

# --- Calendário
ws = wb.create_sheet("Calendário")
cols = ["Data", "Dia da semana", "Rede", "Formato", "Pilar", "Funil", "Objetivo", "Tema", "Gancho (completo)", "Chamada para ação", "Stories do dia (tipo e gancho)", "Status"]
cabecalho(ws, 1, cols)
N = 40
for r in range(2, N + 2):
    ws.cell(row=r, column=2, value=f'=IF(A{r}="","",TEXT(A{r},"dddd"))')
    ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).border = border
        ws.cell(row=r, column=c).alignment = wrap
def lista(ws, col, opcoes):
    dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{N+1}")
lista(ws, "C", ["Instagram", "TikTok", "Facebook", "LinkedIn", "YouTube", "WhatsApp"])
lista(ws, "D", ["Reels", "Carrossel", "Foto", "Stories", "Vídeo longo", "Texto"])
lista(ws, "E", ["Educar", "Conectar", "Provar", "Vender"])
lista(ws, "F", ["Topo", "Meio", "Fundo"])
lista(ws, "G", ["Atrair", "Engajar", "Conectar", "Vender"])
lista(ws, "L", ["A fazer", "Em produção", "Pronto", "Agendado", "Publicado"])
# cor de status
verde = PatternFill("solid", bgColor="CDEAC0"); amarelo = PatternFill("solid", bgColor="FFF2B3")
ws.conditional_formatting.add(f"L2:L{N+1}", CellIsRule(operator="equal", formula=['"Publicado"'], fill=verde))
ws.conditional_formatting.add(f"L2:L{N+1}", CellIsRule(operator="equal", formula=['"Agendado"'], fill=verde))
ws.conditional_formatting.add(f"L2:L{N+1}", CellIsRule(operator="equal", formula=['"Em produção"'], fill=amarelo))
larguras(ws, {"A": 12, "B": 14, "C": 12, "D": 12, "E": 11, "F": 9, "G": 11, "H": 30, "I": 42, "J": 26, "K": 34, "L": 13})
ws.freeze_panes = "C2"
# exemplo na 1ª linha (apague ao começar)
ex = ["", "", "Instagram", "Carrossel", "Educar", "Topo", "Atrair", "(exemplo) Sinais de que a ração não serve", "3 sinais de que a ração do seu cão não está servindo", "Salve para consultar depois", "Enquete: 'Seu cão come rápido demais?'", "A fazer"]
for c, v in enumerate(ex, 1):
    if v: ws.cell(row=2, column=c, value=v)
    ws.cell(row=2, column=c).font = Font(italic=True, color="8A7A52")

# --- Resumo
ws = wb.create_sheet("Resumo")
cabecalho(ws, 1, ["Pilar", "Posts no mês", "% do total", "Meta"])
metas = {"Educar": 0.40, "Conectar": 0.25, "Provar": 0.20, "Vender": 0.15}
for i, (p, m) in enumerate(metas.items(), 2):
    ws.cell(row=i, column=1, value=p)
    ws.cell(row=i, column=2, value=f"=COUNTIF('Calendário'!E2:E{N+1},A{i})")
    ws.cell(row=i, column=3, value=f"=IF(SUM($B$2:$B$5)=0,0,B{i}/SUM($B$2:$B$5))"); ws.cell(row=i, column=3).number_format = "0%"
    ws.cell(row=i, column=4, value=m); ws.cell(row=i, column=4).number_format = "0%"
ws.cell(row=6, column=1, value="Total"); ws.cell(row=6, column=2, value="=SUM(B2:B5)")
cabecalho(ws, 8, ["Formato", "Posts no mês", "", ""])
for i, f in enumerate(["Reels", "Carrossel", "Foto", "Stories"], 9):
    ws.cell(row=i, column=1, value=f)
    ws.cell(row=i, column=2, value=f"=COUNTIF('Calendário'!D2:D{N+1},A{i})")
cabecalho(ws, 14, ["Funil", "Posts no mês", "", ""])
for i, f in enumerate(["Topo", "Meio", "Fundo"], 15):
    ws.cell(row=i, column=1, value=f)
    ws.cell(row=i, column=2, value=f"=COUNTIF('Calendário'!F2:F{N+1},A{i})")
ws.cell(row=19, column=1, value="Confira: nenhum post de venda seguido de outro? Os três momentos do funil aparecem?").font = Font(italic=True)
larguras(ws, {"A": 18, "B": 16, "C": 12, "D": 10})

# --- Datas
ws = wb.create_sheet("Datas")
cabecalho(ws, 1, ["Mês", "Datas fixas e campanhas", "Datas do meu nicho (preencha)"])
datas = [
 ("Janeiro", "Ano-Novo (1/1) · Janeiro Branco (saúde mental) · Volta às aulas e liquidações de verão"),
 ("Fevereiro", "Carnaval (data móvel) · Fevereiro Laranja e Roxo (campanhas de saúde) · Volta às aulas"),
 ("Março", "Dia Internacional da Mulher (8/3) · Dia do Consumidor (15/3) · Início do outono"),
 ("Abril", "Páscoa (móvel) · Tiradentes (21/4) · Abril Azul"),
 ("Maio", "Dia do Trabalhador (1/5) · Dia das Mães (2º domingo)"),
 ("Junho", "Dia do Meio Ambiente (5/6) · Dia dos Namorados (12/6) · Festas juninas"),
 ("Julho", "Dia do Amigo (20/7) · Férias escolares"),
 ("Agosto", "Dia dos Pais (2º domingo) · Dia do Estudante (11/8) · Dia do Folclore (22/8)"),
 ("Setembro", "Independência (7/9) · Setembro Amarelo (10/9) · Dia do Cliente (15/9) · Início da primavera"),
 ("Outubro", "Dia das Crianças (12/10) · Dia do Professor (15/10) · Dia do Comerciante (16/10) · Outubro Rosa"),
 ("Novembro", "Finados (2/11) · Consciência Negra (20/11) · Novembro Azul · Black Friday e Cyber Monday (fim do mês, a data varia)"),
 ("Dezembro", "Dezembro Vermelho · Natal (25/12) · Virada do ano"),
]
for i, (m, d) in enumerate(datas, 2):
    ws.cell(row=i, column=1, value=m); ws.cell(row=i, column=2, value=d)
    for c in range(1, 4):
        ws.cell(row=i, column=c).border = border; ws.cell(row=i, column=c).alignment = wrap
ws.cell(row=15, column=1, value="Confirme as datas todo ano: várias são móveis. Acrescente as datas do seu nicho (ex.: Dia do Gato e Dia do Cachorro para pet shops).").font = Font(italic=True)
larguras(ws, {"A": 14, "B": 90, "C": 45})

# --- Banco de ganchos
ws = wb.create_sheet("Banco de ganchos")
cabecalho(ws, 1, ["Pilar", "Modelo de gancho (troque o que está entre colchetes)", "Meu gancho"])
ganchos = {
 "Educar": ["3 erros que [seu público] comete ao [ação]", "O passo a passo para [resultado] em [tempo]", "Mito ou verdade: [crença comum do seu nicho]", "Se eu começasse [atividade] hoje, eu faria isto", "[Número] sinais de que [problema] está acontecendo", "O que ninguém te conta sobre [tema]", "Como escolher [produto ou serviço]: o que olhar antes de comprar", "[Tema] explicado em 30 segundos", "Antes de [ação], faça estas [número] perguntas", "Diferença entre [A] e [B] (e qual serve para você)"],
 "Conectar": ["O dia em que quase desisti de [coisa] (e o que me fez continuar)", "Bastidores: como [produto] é feito, do começo ao fim", "Um dia comigo no [negócio]", "Por que eu comecei [negócio]", "O que eu aprendi errando em [assunto]", "Uma coisa que eu acredito e poucos no meu mercado dizem", "A história por trás do nome [do negócio ou produto]", "Meu maior erro como [profissão] (e a lição)", "Apresentando a equipe: quem cuida de você aqui", "Uma rotina minha que melhora o meu trabalho"],
 "Provar": ["Antes e depois de [cliente, com autorização]", "O que a [cliente] disse depois de [resultado]", "A pergunta que mais recebo: '[pergunta]'. Respondo aqui", "[Número] clientes atendidos em [tempo]", "Mensagens que me fazem querer continuar", "Como foi o processo do [cliente] do início ao fim", "O resultado de [ação] em [prazo], com números reais", "O que acontece depois que você contrata [serviço]", "Perguntas e objeções: '[objeção]'. Resposta honesta", "Um erro de cliente que a gente evitou para ele"],
 "Vender": ["Para quem [produto] é (e para quem não é)", "O que vem no [produto], item por item", "Como pedir: passo a passo em 3 etapas", "Últimas vagas ou última semana, com data real", "[Produto] resolve [dor]: veja como", "Comparação: [seu produto] versus [alternativa comum]", "Perguntas frequentes sobre [produto]", "Condição especial até [data]", "Responda '[PALAVRA]' e eu te mando os detalhes", "Resultados de quem já comprou [produto]"],
}
r = 2
for pilar, lst in ganchos.items():
    for g in lst:
        ws.cell(row=r, column=1, value=pilar); ws.cell(row=r, column=2, value=g)
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = border; ws.cell(row=r, column=c).alignment = wrap
        r += 1
larguras(ws, {"A": 12, "B": 70, "C": 50}); ws.freeze_panes = "A2"

# --- Revisão do mês
ws = wb.create_sheet("Revisão do mês")
cabecalho(ws, 1, ["Formato", "Nº de posts", "Alcance médio", "Salvamentos", "Compartilhamentos", "Mensagens", "Repetir? (sim/não)", "Observação"])
for i, f in enumerate(["Reels", "Carrossel", "Foto", "Stories"], 2):
    ws.cell(row=i, column=1, value=f)
    for c in range(1, 9): ws.cell(row=i, column=c).border = border
cabecalho(ws, 8, ["Top 3 do mês", "Pilar", "Formato", "Gancho", "Dia e horário", "Por que funcionou?", "Variação para o próximo mês", ""])
for i in range(9, 12):
    ws.cell(row=i, column=1, value=f"{i-8}º")
    for c in range(1, 8): ws.cell(row=i, column=c).border = border
larguras(ws, {"A": 16, "B": 12, "C": 14, "D": 36, "E": 16, "F": 30, "G": 34, "H": 20})

wb.save("planilha-calendario-conteudo.xlsx")
print("ok planilha-calendario-conteudo.xlsx")
