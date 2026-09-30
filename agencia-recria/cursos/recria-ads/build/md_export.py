"""Converte os .md do Recria Ads em PDF (visual Recria) e DOCX editável."""
import pathlib, re, base64, markdown
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from playwright.sync_api import sync_playwright
S=pathlib.Path(__file__).parent
BASE=pathlib.Path('/home/user/recria-clientes/agencia-recria/cursos/recria-ads')
OUT=BASE/'para-ler-e-gravar'; OUT.mkdir(exist_ok=True)
FONTS=(S/'c1/fonts_embedded.css').read_text()
LOGO='data:image/png;base64,'+base64.b64encode((S/'ck/logo_sem_tagline.png').read_bytes()).decode()
GOLD=RGBColor(0x9a,0x74,0x24); BLACK=RGBColor(0x0b,0x0b,0x0b)
CSS=FONTS+"""@page{size:A4;margin:18mm 16mm}
body{font-family:'Lora',serif;font-size:12.5pt;line-height:1.6;color:#1d1d1d}
h1{font-family:'Playfair Display';font-size:24pt;border-bottom:3px solid #C9A24E;padding-bottom:6px;page-break-before:always}
h1:first-of-type{page-break-before:avoid}
h2{font-family:'Playfair Display';font-size:16pt;background:#0b0b0b;color:#E4C988;padding:7px 12px;margin-top:20px;page-break-after:avoid}
h3{font-family:'Playfair Display';font-size:13.5pt;color:#5a4412}
strong{color:#5a4412} table{border-collapse:collapse;width:100%;font-size:10.5pt;margin:8px 0}
th{background:#0b0b0b;color:#E4C988;text-align:left;padding:6px} td{border-bottom:1px solid #e1d5ba;padding:6px;vertical-align:top}
hr{border:none;border-top:1px dashed #C9A24E;margin:16px 0} pre,code{white-space:pre-wrap;font-size:10.5pt;background:#f3eee3}
blockquote{border-left:4px solid #C9A24E;margin:8px 0;padding:4px 14px;background:#faf6ee;font-style:italic}
.capa{height:250mm;display:flex;flex-direction:column;justify-content:center;background:#0b0b0b;color:#fff;padding:0 18mm;margin:-18mm -16mm 0;page-break-after:always}
.capa h1{color:#fff;border:none;font-size:34pt;page-break-before:avoid} .capa p{color:#E4C988;font-size:14pt}"""

def pdf(md, title, sub, out):
    body=markdown.markdown(md,extensions=['tables','fenced_code','nl2br'])
    capa=f'<div class="capa"><img src="{LOGO}" style="height:90px;align-self:flex-start;margin-bottom:40px"><p>RECRIA ADS</p><h1>{title}</h1><p>{sub}</p></div>'
    html=f'<!doctype html><html><head><meta charset="utf-8"><style>{CSS}</style></head><body>{capa}{body}</body></html>'
    with sync_playwright() as p:
        b=p.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome'); pg=b.new_page()
        pg.set_content(html,wait_until='networkidle'); pg.evaluate('document.fonts.ready'); pg.pdf(path=str(out),format='A4',print_background=True); b.close()

def runs(par, text, italic=False):
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not part: continue
        if part.startswith('**'): r=par.add_run(part[2:-2]); r.bold=True; r.font.color.rgb=RGBColor(0x5a,0x44,0x12)
        elif part.startswith('*') and len(part)>1: r=par.add_run(part[1:-1]); r.italic=True
        else: r=par.add_run(part)
        if italic: r.italic=True

def docx(md, title, sub, out):
    d=Document()
    for s in d.sections: s.left_margin=s.right_margin=Cm(2.2); s.top_margin=s.bottom_margin=Cm(2)
    st=d.styles['Normal']; st.font.name='Georgia'; st.font.size=Pt(12)
    t=d.add_paragraph(); r=t.add_run(title); r.bold=True; r.font.size=Pt(26); r.font.color.rgb=BLACK
    s2=d.add_paragraph(); r=s2.add_run('RECRIA ADS · '+sub); r.font.color.rgb=GOLD
    lines=md.split('\n'); i=0; incode=False
    while i<len(lines):
        ln=lines[i].rstrip()
        if ln.startswith('```'): incode=not incode; i+=1; continue
        if incode: p=d.add_paragraph(); r=p.add_run(ln); r.font.name='Courier New'; r.font.size=Pt(10); i+=1; continue
        if ln.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:|-]+\|$',lines[i+1].strip()):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s:|-]+\|$',lines[i].strip()): rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i+=1
            tb=d.add_table(rows=len(rows),cols=len(rows[0])); tb.style='Table Grid'
            for ri,row in enumerate(rows):
                for ci,c in enumerate(row[:len(rows[0])]):
                    cell=tb.cell(ri,ci); cell.text=''; runs(cell.paragraphs[0],c)
                    if ri==0:
                        for r in cell.paragraphs[0].runs: r.bold=True
            d.add_paragraph(); continue
        m=re.match(r'^(#{1,3}) (.*)',ln)
        if m:
            lvl=len(m.group(1)); h=d.add_heading(level=lvl); r=h.add_run(m.group(2)); r.font.color.rgb=BLACK if lvl==1 else GOLD
            if lvl==1 and i>0: pass
        elif ln.strip()=='---': d.add_paragraph('—' * 20).alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r'^\s*[-*] ',ln): p=d.add_paragraph(style='List Bullet'); runs(p,re.sub(r'^\s*[-*] ','',ln))
        elif re.match(r'^\s*\d+\. ',ln): p=d.add_paragraph(style='List Number'); runs(p,re.sub(r'^\s*\d+\. ','',ln))
        elif ln.startswith('>'): p=d.add_paragraph(); runs(p,ln.lstrip('> '),italic=True)
        elif ln.strip(): p=d.add_paragraph(); runs(p,ln)
        i+=1
    d.save(str(out))

R=BASE/'roteiros-completos'
DOCS=[
 ('Roteiros - Modulos 0 a 2', ''.join((R/f).read_text()+'\n\n' for f in ['00-legenda.md','01-modulo-0.md','02-modulo-1.md','03-modulo-2.md']), 'Roteiros de gravação: Módulos 0, 1 e 2'),
 ('Roteiros - Modulos 3 a 7', ''.join((R/f).read_text()+'\n\n' for f in ['00b-legenda-modulos-3-a-7.md','04-modulo-3.md','05-modulo-4.md','06-modulo-5.md','07-modulo-6.md','08-modulo-7.md']), 'Roteiros de gravação: Módulos 3, 4, 5, 6 e 7'),
 ('Video de vendas e pagina de vendas', (BASE/'pagina-de-vendas.md').read_text(), 'Roteiro do vídeo de vendas (VSL) e textos da página'),
 ('Copies dos anuncios', (BASE/'criativos/copies-dos-anuncios.md').read_text(), 'Textos dos anúncios, públicos e roteiros de Reels'),
 ('Guia de pronuncia dos autores', (BASE/'guia-de-pronuncia.md').read_text(), 'Como falar o nome de cada autor'),
 ('Ementa do curso', (BASE/'ementa.md').read_text(), 'Estrutura completa, liberação e preços'),
]
for name, md, sub in DOCS:
    title=name.replace('Modulos','Módulos').replace('Video','Vídeo').replace('pagina','página').replace('anuncios','anúncios').replace('pronuncia','pronúncia').replace('Modulo 6','Módulo 6').replace('calendarios','calendários')
    pdf(md, title, sub, OUT/f'{name}.pdf'); docx(md, title, sub, OUT/f'{name}.docx'); print('ok', name)
